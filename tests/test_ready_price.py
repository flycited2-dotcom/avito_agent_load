"""Контракт источника: ошибки, повторы и неполные снимки не меняют каталог."""
from datetime import datetime, timedelta, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sqlite3

from openpyxl import Workbook, load_workbook
import pytest
import yaml

from avito_bridge.ready_price.catalog import HEADERS, SHEET, InvalidSnapshot, parse, price
from avito_bridge.ready_price.store import connect, export_catalog, import_release, source_freshness_problem

NOW = datetime(2026, 9, 21, 12, tzinfo=timezone.utc)


@pytest.fixture
def policy(tmp_path):
    path = tmp_path / "categories.yaml"
    path.write_text(yaml.safe_dump({"groups": {"Мультиварки": {"bucket": "small", "reason": "target"},
                      "Посуда": {"bucket": "excluded", "reason": "outside_scope"}}}, allow_unicode=True), encoding="utf-8")
    return path


def release(root, *, count=4, prices=None, generated=None, omit=(), group="Мультиварки", fallback=False):
    root.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append([None] * 6)
    sheet.append(HEADERS)
    sheet.append([group])
    for i in range(count):
        if i not in omit:
            sheet.append([i + 1, f"00-{i:05}", "Comfee", f"Мультиварка Comfee CF-MC{i} белая",
                          (prices or {}).get(i, 3993), 999])
    file = root / "price.xlsx"
    book.save(file)
    generated = generated or NOW - timedelta(hours=1)
    manifest = {"schema_version": 1, "source": "telegram-nikita", "file": "price.xlsx",
                "source_sha256": hashlib.sha256(file.read_bytes()).hexdigest(),
                "generated_at": generated.isoformat(), "issue_date": generated.date().isoformat(),
                "snapshot_kind": "partial" if fallback else "full",
                "telegram_route": {"bot_id": 8639233666, "chat_id": "-1001929222037"},
                "provenance": {"producer": "excel-automation.transform", "evidence": "transform_success",
                               "supplier_fallback": fallback, "supplier": {"modified_at": generated.isoformat()}}}
    path = root / "release.json"
    path.write_text(json.dumps(manifest), encoding="utf-8")
    return path


def ingest(path, db, policy, **kwargs):
    return import_release(path, db, policy, current=NOW, minimum_rows=1, **kwargs)


def change_manifest(path, **kwargs):
    data = json.loads(path.read_text())
    data.update(kwargs)
    path.write_text(json.dumps(data))


@pytest.mark.parametrize("source,expected", [(3864, 4058), (3993, 4193), (4573, 4802), (4723, 4960),
                                            (19990, 20990), (Decimal("0.01"), 1), (100, 105)])
def test_price_rule(source, expected):
    assert price(source)[1] == expected


@pytest.mark.parametrize("value", [None, 0, -1, True, "", "NaN", "Infinity", "=A1", "1,00"])
def test_invalid_price(value):
    with pytest.raises(InvalidSnapshot):
        price(value)


def test_full_catalog_and_unknown_quantity(tmp_path, policy):
    path = release(tmp_path / "one")
    db = tmp_path / "catalog.db"
    result = ingest(path, db, policy)
    assert result["status"] == "accepted"
    items = export_catalog(db)
    assert len(items) == 4
    assert all(x["avito_price"] == 4193 and x["quantity"] is None for x in items)
    assert items[0]["row"] == 4 and items[0]["article"] == "00-00000"
    assert all(not x["identity_verified"] and not x["publication_allowed"] for x in items)


def test_repeat_and_price_change_do_not_compound(tmp_path, policy):
    db = tmp_path / "catalog.db"
    first = release(tmp_path / "one", generated=NOW - timedelta(days=1))
    assert not ingest(first, db, policy)["duplicate"]
    assert ingest(first, db, policy)["duplicate"]
    second = release(tmp_path / "two", prices={0: 4000})
    assert ingest(second, db, policy)["status"] == "accepted"
    assert ingest(second, db, policy)["duplicate"]
    items = export_catalog(db)
    assert items[0]["source_price"] == "4000" and items[0]["avito_price"] == 4200
    assert items[1]["avito_price"] == 4193


@pytest.mark.parametrize("kind", ["empty", "invalid_price", "stale", "future", "fallback", "checksum", "wrong_source", "out_of_order", "count_drop"])
def test_bad_snapshot_preserves_previous_catalog(tmp_path, policy, kind):
    db = tmp_path / "catalog.db"
    ingest(release(tmp_path / "one", generated=NOW - timedelta(days=1)), db, policy)
    before = export_catalog(db)
    kwargs = {"prices": {0: 1000}}
    if kind == "empty": kwargs["count"] = 0
    if kind == "invalid_price": kwargs["prices"] = {0: "=A1"}
    if kind == "stale": kwargs["generated"] = NOW - timedelta(days=5)
    if kind == "future": kwargs["generated"] = NOW + timedelta(hours=2)
    if kind == "fallback": kwargs["fallback"] = True
    if kind == "out_of_order": kwargs["generated"] = NOW - timedelta(days=2)
    if kind == "count_drop": kwargs["count"] = 2
    path = release(tmp_path / "two", **kwargs)
    if kind == "checksum": change_manifest(path, source_sha256="0" * 64)
    if kind == "wrong_source": change_manifest(path, source="itp")
    result = ingest(path, db, policy)
    assert result["status"] == "quarantined", result
    assert result["reason"]
    assert export_catalog(db) == before


def test_new_export_of_old_supplier_is_stale(tmp_path, policy):
    path = release(tmp_path / "one")
    data = json.loads(path.read_text())
    data["provenance"]["supplier"]["modified_at"] = (NOW - timedelta(days=20)).isoformat()
    path.write_text(json.dumps(data))
    assert ingest(path, tmp_path / "db", policy)["reason"] == "stale_release"


def test_supplier_freshness_checks_original_time_at_publication():
    release_row = {"generated_at": NOW.isoformat(), "manifest": json.dumps({
        "schema_version": 2, "snapshot_kind": "full", "provenance": {
            "supplier_fallback": False, "supplier": {"modified_at": (NOW - timedelta(days=5)).isoformat()}}})}
    assert source_freshness_problem(release_row, current=NOW) == "supplier_snapshot_stale"


def test_missing_item_recorded_without_automatic_delisting(tmp_path, policy):
    db = tmp_path / "catalog.db"
    ingest(release(tmp_path / "one", generated=NOW - timedelta(days=1)), db, policy)
    assert ingest(release(tmp_path / "two", omit=(1,)), db, policy)["status"] == "accepted"
    missing = next(x for x in export_catalog(db) if x["article"] == "00-00001")
    assert missing["present"] is False and missing["quantity"] is None
    assert missing["status"] == "awaiting_absence_policy" and not missing["publication_allowed"]
    assert len(export_catalog(db)) == 4


def test_delivery_can_finish_after_import(tmp_path, policy):
    db = tmp_path / "catalog.db"
    path = release(tmp_path / "one")
    ingest(path, db, policy)
    (path.parent / "telegram.json").write_text('{"ever_sent":true}')
    assert ingest(path, db, policy)["duplicate"]
    with sqlite3.connect(db) as connection:
        assert json.loads(connection.execute("SELECT telegram FROM releases").fetchone()[0])["ever_sent"]


def test_unknown_group_is_visible(tmp_path, policy):
    db = tmp_path / "catalog.db"
    path = release(tmp_path / "one", group="Новый вид техники")
    ingest(path, db, policy)
    assert all(x["status"] == "awaiting_category_review" and x["reason"] == "unknown_group" for x in export_catalog(db))


@pytest.mark.parametrize("mutation,reason", [("duplicate", "duplicate_article"), ("headers", "unexpected_headers"), ("sheet", "unexpected_sheets")])
def test_structure_and_duplicate_articles_fail_closed(tmp_path, mutation, reason):
    path = release(tmp_path / "one").parent / "price.xlsx"
    book = load_workbook(path)
    if mutation == "duplicate": book.active["B5"] = "00-00000"
    if mutation == "headers": book.active["E2"] = "Закупочная цена"
    if mutation == "sheet": book.active.title = "Закупка"
    book.save(path)
    with pytest.raises(InvalidSnapshot, match=reason):
        parse(path, {})


def test_transaction_failure_rolls_back_catalog_and_release(tmp_path, policy):
    db = tmp_path / "catalog.db"
    first = release(tmp_path / "one", generated=NOW - timedelta(days=1))
    ingest(first, db, policy)
    before = export_catalog(db)
    connection = connect(db)
    connection.execute("CREATE TRIGGER fail_catalog BEFORE UPDATE ON catalog BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    connection.close()
    with pytest.raises(sqlite3.IntegrityError):
        ingest(release(tmp_path / "two", prices={0: 1000}), db, policy)
    assert export_catalog(db) == before
    with sqlite3.connect(db) as connection:
        assert connection.execute("SELECT COUNT(*) FROM releases").fetchone()[0] == 1


def test_supplier_workbook_is_stock_authority(tmp_path, policy):
    path = release(tmp_path / "one")
    raw = Workbook()
    sheet = raw.active
    for _ in range(3):
        sheet.append([None] * 8)
    for i in range(500):
        sheet.append([f"00-{i:05}", "group", "Comfee", f"Item {i}", 100, "Под заказ"])
    # Финальная книга содержит артикул, которого нет в исходном полном прайсе.
    sheet["A4"] = "different-article"
    raw_path = path.parent / "supplier.xlsx"
    raw.save(raw_path)
    manifest = json.loads(path.read_text())
    manifest["schema_version"] = 2
    manifest["supplier_file"] = "supplier.xlsx"
    manifest["provenance"]["supplier"]["sha256"] = hashlib.sha256(raw_path.read_bytes()).hexdigest()
    path.write_text(json.dumps(manifest))
    result = ingest(path, tmp_path / "catalog.db", policy)
    assert result["status"] == "accepted", result
    items = {x["article"]: x for x in export_catalog(tmp_path / "catalog.db")}
    assert items["00-00000"]["availability"] == "unverified_origin"
    assert items["00-00001"]["availability"] == "supplier_price_present"
    raw_path.write_bytes(b"corrupt")
    assert ingest(path, tmp_path / "other.db", policy)["reason"] == "supplier_checksum_mismatch"
