"""Read the current mail price list without confusing image IDs with prices.

Missing articles have UNKNOWN availability, not zero. This reader never
changes publication or prices; warehouse names must be mapped separately.
"""
from datetime import date
from pathlib import Path
import re

from openpyxl import load_workbook

from avito_bridge.ingest.price_xls import MAX_WORKBOOK_BYTES, MAX_WORKBOOK_ROWS

MONTHS = 'января февраля марта апреля мая июня июля августа сентября октября ноября декабря'.split()


def availability(value):
    if value is None or str(value).strip() == '':
        return None
    text = str(value).strip().casefold()
    if text == 'в наличии':
        return True
    if text in {'нет в наличии', 'нет', 'под заказ'}:
        return False
    try:
        number = float(text.replace('\xa0', '').replace(' ', '').replace(',', '.'))
    except ValueError:
        raise ValueError(f'Unknown stock label: {value!r}') from None
    if not 0 <= number < float('inf'):
        raise ValueError('Invalid stock quantity')
    return number > 0


def read_availability(path: Path, *, today=None, max_age_days=2):
    today = today or date.today()
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError('Workbook too large')
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        if sheet.max_row > MAX_WORKBOOK_ROWS:
            raise ValueError('Too many rows')
        rows = list(sheet.iter_rows(values_only=True))
        heading = ' '.join(str(v) for row in rows[:2] for v in row if v)
        match = re.search(r'Прайс-лист на\s+(\d+)\s+(\w+)\s+(\d{4})', heading)
        if not match or match[2] not in MONTHS:
            raise ValueError('Missing price-list date')
        published = date(int(match[3]), MONTHS.index(match[2]) + 1, int(match[1]))
        if not 0 <= (today - published).days <= max_age_days:
            raise ValueError('Stale or future price list')
        header = rows[2]
        expected = {0: 'Артикул', 4: 'Номенклатура', 13: 'Изображение',
                    14: 'от 100 т.руб/мес', 15: 'Склад основной ООО', 16: 'Склад управленки'}
        if any(len(header) <= i or str(header[i]).strip() != label for i, label in expected.items()):
            raise ValueError('Unsupported supplier schema')
        result = {}
        for row in rows[4:]:
            if len(row) < 17 or row[14] is None:
                continue  # Group headers may also contain an article and stock label.
            try:
                price = float(str(row[14]).replace('\xa0', '').replace(' ', '').replace(',', '.'))
            except ValueError:
                raise ValueError('Invalid product price') from None
            if not 0 < price < float('inf') or not row[0] or not row[4]:
                raise ValueError('Invalid product row')
            article = str(row[0]).strip()
            if article in result:
                raise ValueError('Duplicate supplier article')
            result[article] = {'name': str(row[4]).strip(),
                               'main_available': availability(row[15]),
                               'management_available': availability(row[16])}
        if not result:
            raise ValueError('Empty supplier workbook')
        return {'date': published.isoformat(), 'items': result}
    finally:
        book.close()
