"""Транзакционный каталог; публикация и снятие объявлений здесь не выполняются."""
from __future__ import annotations

from dataclasses import asdict, replace
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import sqlite3
import xlrd
import openpyxl

from .catalog import InvalidSnapshot, load_categories, parse, summary

SOURCE = "telegram-nikita"


def now_utc():
    return datetime.now(timezone.utc)


def timestamp(value):
    result = datetime.fromisoformat(value)
    if result.tzinfo is None:
        raise InvalidSnapshot("timezone_required")
    return result.astimezone(timezone.utc)


def connect(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path, timeout=30)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA foreign_keys=ON")
    db.executescript("""
        CREATE TABLE IF NOT EXISTS releases (
            sha256 TEXT PRIMARY KEY, source TEXT NOT NULL, generated_at TEXT,
            imported_at TEXT NOT NULL, status TEXT NOT NULL, reason TEXT NOT NULL,
            manifest TEXT NOT NULL, summary TEXT NOT NULL, telegram TEXT NOT NULL DEFAULT '{}'
        );
        CREATE TABLE IF NOT EXISTS snapshot_items (
            sha256 TEXT NOT NULL REFERENCES releases(sha256), article TEXT NOT NULL,
            row_number INTEGER NOT NULL, data TEXT NOT NULL, PRIMARY KEY(sha256, article)
        );
        CREATE TABLE IF NOT EXISTS catalog (
            source TEXT NOT NULL, article TEXT NOT NULL, sha256 TEXT NOT NULL REFERENCES releases(sha256),
            data TEXT NOT NULL, present INTEGER NOT NULL, missing_since TEXT,
            PRIMARY KEY(source, article)
        );
        CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY, created_at TEXT NOT NULL, kind TEXT NOT NULL, detail TEXT NOT NULL
        );
    """)
    return db


def validate_release(manifest: dict, file: Path, current: datetime, max_age_days: int):
    if manifest.get("schema_version") not in {1, 2} or manifest.get("source") != SOURCE:
        raise InvalidSnapshot("unexpected_source_or_schema")
    if manifest.get("file") != "price.xlsx":
        raise InvalidSnapshot("unexpected_snapshot_path")
    digest = hashlib.sha256(file.read_bytes()).hexdigest()
    if digest != manifest.get("source_sha256"):
        raise InvalidSnapshot("checksum_mismatch")
    provenance = manifest.get("provenance") or {}
    if provenance.get("producer") != "excel-automation.transform" or provenance.get("evidence") not in {
        "transform_success", "verified_daily_log"
    }:
        raise InvalidSnapshot("unverified_provenance")
    if manifest.get("telegram_route") != {"bot_id": 8639233666, "chat_id": "-1001929222037"}:
        raise InvalidSnapshot("unexpected_source_route")
    generated = timestamp(manifest["generated_at"])
    issue = date.fromisoformat(manifest["issue_date"])
    supplier_time = timestamp(provenance["supplier"]["modified_at"])
    if generated > current + timedelta(minutes=10) or issue > current.date() + timedelta(days=1):
        raise InvalidSnapshot("future_release")
    # Возраст самого прайса И его исходника; новый экспорт старого источника не омолаживает наличие.
    oldest = min(generated, supplier_time, datetime.combine(issue, datetime.min.time(), timezone.utc))
    if current - oldest > timedelta(days=max_age_days):
        raise InvalidSnapshot("stale_release")
    if manifest.get("snapshot_kind") != "full" or provenance.get("supplier_fallback") is not False:
        raise InvalidSnapshot("partial_or_fallback_snapshot")
    return digest, generated


def supplier_articles(manifest: dict, directory: Path) -> set[str] | None:
    """Артикул в свежем прайсе означает наличие; «Под заказ» у этого поставщика допустимо."""
    if manifest.get("schema_version") == 1:
        return None  # старый выпуск не содержит проверяемого исходного прайса
    name = manifest.get("supplier_file")
    if name not in {"supplier.xls", "supplier.xlsx"}:
        raise InvalidSnapshot("missing_supplier_workbook")
    path = directory / name
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise InvalidSnapshot("missing_supplier_workbook") from exc
    if hashlib.sha256(data).hexdigest() != manifest["provenance"]["supplier"]["sha256"]:
        raise InvalidSnapshot("supplier_checksum_mismatch")
    try:
        if name.endswith(".xls"):
            book = xlrd.open_workbook(file_contents=data)
            if book.nsheets != 1:
                raise InvalidSnapshot("unexpected_supplier_sheets")
            sheet = book.sheet_by_index(0)
            rows = (sheet.row_values(i) for i in range(3, sheet.nrows))
        else:
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if len(book.sheetnames) != 1:
                raise InvalidSnapshot("unexpected_supplier_sheets")
            rows = book.active.iter_rows(min_row=4, values_only=True)
        articles = set()
        for row in rows:
            if len(row) < 5 or not row[0] or not row[3] or not isinstance(row[4], (int, float)) or row[4] <= 0:
                continue
            article = str(row[0]).strip().casefold()
            if article in articles:
                raise InvalidSnapshot("duplicate_supplier_article")
            articles.add(article)
        if len(articles) < 500:
            raise InvalidSnapshot("supplier_row_count_below_minimum")
        return articles
    except (xlrd.XLRDError, OSError, ValueError, TypeError) as exc:
        if isinstance(exc, InvalidSnapshot):
            raise
        raise InvalidSnapshot("unreadable_supplier_workbook") from exc
    finally:
        if name.endswith(".xlsx") and "book" in locals():
            book.close()


def import_release(manifest_path: Path, database: Path, categories: Path, *, current=None,
                   max_age_days=4, minimum_rows=500, max_count_change=0.25):
    """Один SHA импортируется один раз. Вся книга проверяется до изменения каталога."""
    current = current or now_utc()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    file = manifest_path.parent / "price.xlsx"
    # Ключ БД — фактические байты, не недоверенное поле manifest.
    digest = hashlib.sha256(file.read_bytes()).hexdigest()
    generated = str(manifest.get("generated_at") or "")
    status, reason, items, supplier = "accepted", "", [], None
    try:
        _, generated_dt = validate_release(manifest, file, current, max_age_days)
        generated = generated_dt.isoformat()
        items = parse(file, load_categories(categories))
        supplier = supplier_articles(manifest, manifest_path.parent)
        items = [replace(item, availability=("supplier_price_present" if item.article.casefold() in supplier
                       else "unverified_origin") if supplier is not None else "unverified_origin") for item in items]
        if len(items) < minimum_rows:
            raise InvalidSnapshot("row_count_below_minimum")
    except (InvalidSnapshot, KeyError, TypeError, ValueError) as exc:
        status = "quarantined"
        reason = str(exc) if isinstance(exc, InvalidSnapshot) else "invalid_release_metadata"
    stats = summary(items)
    if status == "accepted" and supplier is not None:
        stats["supplier_rows"] = len(supplier)
    delivery_file = manifest_path.parent / "telegram.json"
    delivery = json.loads(delivery_file.read_text(encoding="utf-8")) if delivery_file.exists() else {}
    db = connect(database)
    try:
        db.execute("BEGIN IMMEDIATE")
        previous = db.execute("SELECT status,reason,summary FROM releases WHERE sha256=?", (digest,)).fetchone()
        if previous:
            # Доставка Telegram может завершиться после первого импорта.
            db.execute("UPDATE releases SET telegram=? WHERE sha256=?", (json.dumps(delivery), digest))
            db.commit()
            return {"sha256": digest, "status": previous["status"], "reason": previous["reason"],
                    "duplicate": True, "summary": json.loads(previous["summary"])}
        latest = db.execute("SELECT * FROM releases WHERE source=? AND status='accepted' ORDER BY generated_at DESC LIMIT 1", (SOURCE,)).fetchone()
        if status == "accepted" and latest:
            if generated <= latest["generated_at"]:
                status, reason = "quarantined", "out_of_order_release"
            else:
                old_stats = json.loads(latest["summary"])
                previous_count = old_stats["rows"]
                if abs(len(items) / previous_count - 1) > max_count_change:
                    status, reason = "quarantined", "abrupt_row_count_change"
                if old_stats.get("supplier_rows") and stats.get("supplier_rows"):
                    if abs(stats["supplier_rows"] / old_stats["supplier_rows"] - 1) > max_count_change:
                        status, reason = "quarantined", "abrupt_supplier_count_change"
                # Резкое обнуление целевого раздела задерживаем; обычный рост
                # ассортимента проверяется общим размером и исходным файлом.
                for bucket in ("small", "large", "tv", "climate"):
                    old_n = old_stats["buckets"].get(bucket, 0)
                    new_n = stats["buckets"].get(bucket, 0)
                    if old_n >= 50 and new_n < old_n / 2:
                        status, reason = "quarantined", f"abrupt_bucket_count_change:{bucket}"
        db.execute("INSERT INTO releases VALUES (?,?,?,?,?,?,?,?,?)",
                   (digest, SOURCE, generated, current.isoformat(), status, reason,
                    json.dumps(manifest, ensure_ascii=False), json.dumps(stats, ensure_ascii=False), json.dumps(delivery),))
        # Snapshot rows сохраняются и для понятного разбора карантина по числу товаров.
        db.executemany("INSERT INTO snapshot_items VALUES (?,?,?,?)",
                       [(digest, x.article, x.row, json.dumps(asdict(x), ensure_ascii=False)) for x in items])
        if status == "accepted":
            db.execute("UPDATE catalog SET present=0, missing_since=COALESCE(missing_since,?) WHERE source=?",
                       (generated, SOURCE))
            db.executemany("""INSERT INTO catalog VALUES (?,?,?,?,1,NULL)
                ON CONFLICT(source,article) DO UPDATE SET sha256=excluded.sha256,data=excluded.data,
                    present=1,missing_since=NULL""",
                [(SOURCE, x.article, digest, json.dumps(asdict(x), ensure_ascii=False)) for x in items])
        db.execute("INSERT INTO events(created_at,kind,detail) VALUES (?,?,?)",
                   (current.isoformat(), status, json.dumps({"sha256": digest, "reason": reason})))
        db.commit()
        return {"sha256": digest, "status": status, "reason": reason, "duplicate": False, "summary": stats}
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def export_catalog(database: Path):
    db = connect(database)
    try:
        result = []
        for row in db.execute("SELECT * FROM catalog ORDER BY article"):
            item = json.loads(row["data"])
            item.update(source=row["source"], source_sha256=row["sha256"], present=bool(row["present"]),
                        missing_since=row["missing_since"], publication_allowed=False)
            if not item["present"]:
                item.update(availability="absent_from_latest_full_snapshot", status="awaiting_absence_policy",
                            reason="absence_recorded_no_automatic_delisting_at_import_stage")
            result.append(item)
        return result
    finally:
        db.close()
