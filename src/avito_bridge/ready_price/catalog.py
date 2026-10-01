"""Строгое чтение итоговой книги. Ошибка снимка никогда не означает пустой склад."""
from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path
import re
import unicodedata

import openpyxl
import yaml

SHEET = "Прайс клиента"
HEADERS = ("№", "Артикул", "Бренд", "Наименование", "Цена (руб.)", "Заказ (шт.)")


class InvalidSnapshot(ValueError):
    pass


def clean(value: object) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def price(value: object) -> tuple[str, int]:
    if isinstance(value, bool) or value is None:
        raise InvalidSnapshot("empty_or_invalid_price")
    try:
        amount = Decimal(str(value))
    except InvalidOperation as exc:
        raise InvalidSnapshot("empty_or_invalid_price") from exc
    if not amount.is_finite() or amount <= 0 or amount > 100_000_000:
        raise InvalidSnapshot("empty_or_invalid_price")
    return str(amount), int((amount * Decimal("1.05")).to_integral_value(rounding=ROUND_CEILING))


def model_hint(name: str, brand: str) -> str:
    """Только поисковая подсказка, НЕ доказательство идентичности модели/цвета."""
    tail = name
    if brand:
        match = re.search(re.escape(brand), name, re.IGNORECASE)
        if match:
            tail = name[match.end():].strip()
    tokens = tail.split()
    for index, token in enumerate(tokens):
        if re.search(r"[A-Za-zА-Яа-я]", token) and re.search(r"\d", token):
            if not re.search(r"(?:Вт|кВт|л|см|BTU)$", token, re.IGNORECASE):
                # Пробелы в модели сохраняем для последующей проверки по источнику.
                start = index - 1 if index and re.fullmatch(r"[A-Z]{1,4}", tokens[index - 1]) else index
                return " ".join(tokens[start:index + 1]).strip(",()")
    return ""


@dataclass(frozen=True)
class Item:
    row: int
    article: str
    brand: str
    name: str
    group: str
    bucket: str
    source_price: str
    avito_price: int
    model_hint: str
    status: str
    reason: str
    sheet: str = SHEET
    quantity: None = None
    availability: str = "in_stock"
    identity_verified: bool = False


def load_categories(path: Path) -> dict[str, dict]:
    policy = yaml.safe_load(path.read_text(encoding="utf-8"))
    groups = policy["groups"]
    allowed = {"small", "large", "tv", "climate", "excluded", "review"}
    if not isinstance(groups, dict) or any(x.get("bucket") not in allowed or not x.get("reason") for x in groups.values()):
        raise ValueError("Invalid category policy")
    return groups


def parse(path: Path, groups: dict[str, dict]) -> list[Item]:
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidSnapshot("unreadable_workbook") from exc
    try:
        if book.sheetnames != [SHEET]:
            raise InvalidSnapshot("unexpected_sheets")
        sheet = book[SHEET]
        if not sheet.max_row or sheet.max_row > 100_000 or sheet.max_column != 6:
            raise InvalidSnapshot("unexpected_dimensions")
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)  # баннер
        if tuple(clean(x) for x in next(rows, ())) != tuple(clean(x) for x in HEADERS):
            raise InvalidSnapshot("unexpected_headers")
        group = ""
        items: list[Item] = []
        seen: set[str] = set()
        errors: list[str] = []
        for number, values in enumerate(rows, 3):
            index, article, brand, name, amount, _order = values
            if all(x is None or x == "" for x in values):
                continue
            if isinstance(index, str) and clean(index) and all(x is None or x == "" for x in values[1:]):
                group = clean(index)
                continue
            sku = clean(article)
            try:
                if not sku or not clean(name) or not group:
                    raise InvalidSnapshot("missing_article_name_or_group")
                if sku.casefold() in seen:
                    raise InvalidSnapshot("duplicate_article")
                seen.add(sku.casefold())
                source_price, avito_price = price(amount)
            except InvalidSnapshot as exc:
                errors.append(f"row={number}:{exc}")
                continue
            rule = groups.get(group, {"bucket": "review", "reason": "unknown_group"})
            bucket = rule["bucket"]
            status = "excluded" if bucket == "excluded" else "awaiting_identity"
            reason = rule["reason"]
            if bucket == "review":
                status = "awaiting_category_review"
            elif bucket == "climate":
                status = "awaiting_source_match"
            elif bucket in {"small", "large", "tv"}:
                reason = "exact_model_variant_and_existing_ad_match_required"
            if not clean(brand) and status == "awaiting_identity":
                reason = "brand_and_exact_model_required"
            items.append(Item(number, sku, clean(brand), clean(name), group, bucket,
                              source_price, avito_price, model_hint(clean(name), clean(brand)), status, reason))
        if errors:
            raise InvalidSnapshot(";".join(errors[:30]) + f";invalid_rows={len(errors)}")
        if not items:
            raise InvalidSnapshot("empty_snapshot")
        return items
    finally:
        book.close()


def summary(items: list[Item]) -> dict:
    return {"rows": len(items), "buckets": dict(Counter(x.bucket for x in items)),
            "statuses": dict(Counter(x.status for x in items)),
            "groups": dict(Counter(x.group for x in items))}


def records(items: list[Item]) -> list[dict]:
    return [asdict(x) for x in items]
