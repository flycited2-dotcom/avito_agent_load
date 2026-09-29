from pathlib import Path
from openpyxl import load_workbook
import json

p = Path('/opt/content-factory/state/prices/mail__info_simfer_com_ru_прайс_лист_опт_xlsx.xlsx')
for data_only in (True,):
    w = load_workbook(p, read_only=True, data_only=data_only)
    print('data_only', data_only)
    for row in w.worksheets[0].iter_rows():
        if row[0].value and row[4].value:
            print(json.dumps([row[i].value for i in (0,4,14,15,16)], ensure_ascii=False, default=str))
    w.close()
