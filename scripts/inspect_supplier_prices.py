"""Read-only workbook source inventory; bounded header/availability inspection."""
import json
from datetime import datetime, timezone
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

folder = Path('/opt/content-factory/state/prices')
candidates = sorted([p for p in folder.glob('*.xlsx') if 'simfer' in p.name.lower() or 'быттехопт' in p.name.lower()], key=lambda p:p.stat().st_mtime, reverse=True)
for p in candidates[:8]:
    print('file',p.name,'mtime',datetime.fromtimestamp(p.stat().st_mtime, timezone.utc).isoformat(),'size',p.stat().st_size)
for p in [Path('/opt/avito-bridge/runtime/appliances/current.xlsx')] + candidates[:1]:
    wb=load_workbook(p,read_only=True,data_only=True)
    print('workbook',str(p),'sheets',wb.sheetnames)
    ws=wb.worksheets[0]
    print('dimensions',ws.max_row,ws.max_column)
    rows=list(ws.iter_rows(max_row=min(ws.max_row or 50000,50000),max_col=10,values_only=True))
    print('headers',json.dumps(rows[:12],ensure_ascii=False,default=str))
    availability=Counter()
    matches=[]
    for row in rows:
        for cell in row:
            if isinstance(cell,str) and any(t in cell.lower() for t in ('под заказ','в наличии','склад')):
                availability[cell]+=1
        if any(any(m in str(cell).upper() for m in ('MS2042DB','MS2044V','MS120W','VT-8513','GL 2162','CT-1658')) for cell in row):
            matches.append(row)
    print('availability_labels',json.dumps(availability.most_common(15),ensure_ascii=False))
    print('selected_models',json.dumps(matches,ensure_ascii=False,default=str))
    wb.close()
